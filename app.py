import io
import zipfile
from pathlib import Path

import pandas as pd
import streamlit as st
import fitz
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from docx import Document
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak

st.set_page_config(page_title="FileMorph AI", page_icon="🔄", layout="wide")

TOOLS = {
    "PDF → PowerPoint": ["pdf"],
    "PowerPoint → PDF": ["pptx"],
    "PDF → Word": ["pdf"],
    "Word → PDF": ["docx"],
    "Excel → CSV": ["xlsx", "xls"],
    "Image → PDF": ["png", "jpg", "jpeg", "webp", "bmp", "tif", "tiff"],
    "PDF → Images": ["pdf"],
    "CSV → PDF": ["csv"],
    "Excel → PDF": ["xlsx", "xls"],
}

st.markdown("""
<style>
.stApp {background:linear-gradient(135deg,#081120,#101b31 55%,#081120);color:#f4f7ff}
section[data-testid="stSidebar"] {background:#0c1729;border-right:1px solid #263653}
.hero {padding:1.25rem 1.5rem;border:1px solid #263c60;border-radius:18px;
background:linear-gradient(120deg,rgba(57,87,180,.28),rgba(0,210,230,.10));margin-bottom:1rem}
.hero h1 {margin:0;font-size:2.25rem}
.hero p {margin:.35rem 0 0;color:#c3d1ed}
.panel {padding:1rem 1.2rem;border:1px solid #2a3b59;border-radius:15px;
background:rgba(20,32,54,.78);margin:.4rem 0 1rem}
div.stButton>button[kind="primary"] {background:linear-gradient(90deg,#6045ff,#00c9df);color:white;border:0;font-weight:700}
</style>
""", unsafe_allow_html=True)

def make_pdf(title, rows, landscape_mode=False, max_rows=500):
    pagesize = landscape(A4) if landscape_mode else A4
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=pagesize, rightMargin=10*mm, leftMargin=10*mm,
                            topMargin=12*mm, bottomMargin=12*mm)
    styles = getSampleStyleSheet()
    story = [Paragraph(str(title).replace("&", "&amp;"), styles["Title"]), Spacer(1, 8)]
    rows = rows[:max_rows + 1]
    if not rows:
        rows = [["No content"]]
    col_count = max(1, max(len(r) for r in rows))
    width = pagesize[0] - 20*mm
    table = Table(rows, repeatRows=1, colWidths=[width / col_count] * col_count)
    table.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#243b64")),
        ("TEXTCOLOR",(0,0),(-1,0),colors.white),
        ("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),
        ("FONTSIZE",(0,0),(-1,-1),7),
        ("GRID",(0,0),(-1,-1),.35,colors.HexColor("#cbd5e1")),
        ("VALIGN",(0,0),(-1,-1),"TOP"),
        ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#eef3f9")]),
        ("LEFTPADDING",(0,0),(-1,-1),4),("RIGHTPADDING",(0,0),(-1,-1),4),
        ("TOPPADDING",(0,0),(-1,-1),4),("BOTTOMPADDING",(0,0),(-1,-1),4),
    ]))
    story.append(table)
    if len(rows) > max_rows:
        story.append(Paragraph(f"Only the first {max_rows} data rows are included.", styles["Italic"]))
    doc.build(story)
    return buffer.getvalue()

def pdf_to_pptx(data, scale):
    pdf = fitz.open(stream=data, filetype="pdf")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    for page in pdf:
        pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(io.BytesIO(pix.tobytes("png")), 0, 0,
                                 width=prs.slide_width, height=prs.slide_height)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()

def pdf_to_docx(data):
    # Cloud-compatible basic text extraction; complex layouts and images are not preserved.
    pdf = fitz.open(stream=data, filetype="pdf")
    doc = Document()
    for i, page in enumerate(pdf):
        if i:
            doc.add_page_break()
        text = page.get_text("text").strip()
        doc.add_paragraph(text if text else "[No selectable text on this page. Scanned pages may require OCR.]")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

def docx_to_pdf(data, title):
    docx = Document(io.BytesIO(data))
    rows = [["Word document text"]]
    for p in docx.paragraphs:
        if p.text.strip():
            rows.append([p.text.strip()])
    for table in docx.tables:
        for row in table.rows:
            rows.append([" | ".join(cell.text.replace("\n", " ") for cell in row.cells)])
    return make_pdf(title, rows)

def pptx_to_pdf(data, title):
    prs = Presentation(io.BytesIO(data))
    rows = [["Slide", "Extracted slide text"]]
    for idx, slide in enumerate(prs.slides, start=1):
        bits = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                bits.append(shape.text.strip().replace("\n", " / "))
        rows.append([str(idx), " | ".join(bits) if bits else "[No text found on this slide]"])
    return make_pdf(title, rows, landscape_mode=True)

def excel_to_csv(data, filename):
    xls = pd.ExcelFile(io.BytesIO(data))
    results = {}
    for sheet in xls.sheet_names:
        df = pd.read_excel(xls, sheet_name=sheet)
        safe = "".join(c if c.isalnum() or c in "-_" else "_" for c in sheet) or "sheet"
        results[f"{Path(filename).stem}_{safe}.csv"] = df.to_csv(index=False).encode("utf-8-sig")
    if len(results) == 1:
        name, payload = next(iter(results.items()))
        return payload, name, "text/csv"
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, payload in results.items():
            z.writestr(name, payload)
    return buf.getvalue(), f"{Path(filename).stem}_csv_sheets.zip", "application/zip"

def image_to_pdf(data):
    im = Image.open(io.BytesIO(data))
    if im.mode != "RGB":
        if im.mode == "RGBA":
            bg = Image.new("RGB", im.size, "white")
            bg.paste(im, mask=im.getchannel("A"))
            im = bg
        else:
            im = im.convert("RGB")
    buf = io.BytesIO()
    im.save(buf, format="PDF", resolution=150)
    return buf.getvalue()

def pdf_to_images(data, image_format, dpi):
    pdf = fitz.open(stream=data, filetype="pdf")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for i, page in enumerate(pdf, 1):
            pix = page.get_pixmap(matrix=fitz.Matrix(dpi/72, dpi/72), alpha=False)
            fmt = image_format.lower()
            z.writestr(f"page_{i:03d}.{fmt}", pix.tobytes(fmt))
    return buf.getvalue()

def csv_to_pdf(data, title):
    df = pd.read_csv(io.BytesIO(data)).fillna("").astype(str)
    rows = [list(map(str, df.columns))] + df.head(500).values.tolist()
    return make_pdf(title, rows, landscape_mode=len(df.columns) > 5, max_rows=500)

def excel_to_pdf(data, title):
    xls = pd.ExcelFile(io.BytesIO(data))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for sheet in xls.sheet_names:
            df = pd.read_excel(xls, sheet_name=sheet).fillna("").astype(str)
            rows = [[str(c) for c in df.columns]] + df.head(500).values.tolist()
            z.writestr(f"{Path(title).stem}_{sheet}.pdf",
                       make_pdf(f"{title} — {sheet}", rows, landscape_mode=len(df.columns)>5))
    # Always ZIP so multi-sheet workbooks can produce one PDF per sheet.
    return buf.getvalue(), f"{Path(title).stem}_pdf_sheets.zip", "application/zip"

def convert(tool, uploaded, quality, dpi, img_fmt):
    data, name = uploaded.getvalue(), uploaded.name
    stem = Path(name).stem
    if tool == "PDF → PowerPoint":
        return pdf_to_pptx(data, quality), f"{stem}.pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    if tool == "PowerPoint → PDF":
        return pptx_to_pdf(data, name), f"{stem}.pdf", "application/pdf"
    if tool == "PDF → Word":
        return pdf_to_docx(data), f"{stem}.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    if tool == "Word → PDF":
        return docx_to_pdf(data, name), f"{stem}.pdf", "application/pdf"
    if tool == "Excel → CSV":
        return excel_to_csv(data, name)
    if tool == "Image → PDF":
        return image_to_pdf(data), f"{stem}.pdf", "application/pdf"
    if tool == "PDF → Images":
        return pdf_to_images(data, img_fmt, dpi), f"{stem}_images.zip", "application/zip"
    if tool == "CSV → PDF":
        return csv_to_pdf(data, name), f"{stem}.pdf", "application/pdf"
    if tool == "Excel → PDF":
        return excel_to_pdf(data, name), f"{stem}_pdf_sheets.zip", "application/zip"
    raise ValueError("Unsupported conversion.")

if "results" not in st.session_state:
    st.session_state.results = []
if "history" not in st.session_state:
    st.session_state.history = []

with st.sidebar:
    st.markdown("## 🔄 FileMorph AI")
    st.caption("All-in-one file converter")
    tool = st.radio("CONVERSION TOOLS", list(TOOLS.keys()))
    st.markdown("---")
    batch = st.toggle("Batch conversion", value=True)
    show_details = st.toggle("Show file details", value=True)
    st.caption("Built for direct Streamlit hosting; no API key required.")

st.markdown("""
<div class="hero">
<h1>🔄 FileMorph <span style="color:#29c8e8">AI</span></h1>
<p>Convert documents, spreadsheets, presentations and images in one workspace.</p>
</div>
""", unsafe_allow_html=True)

main, side = st.columns([2.15, 1])
with main:
    st.subheader(tool)
    notes = {
        "PDF → PowerPoint": "Creates one image-based slide per PDF page. Slide content is not individually editable.",
        "PowerPoint → PDF": "Creates a text-based PDF from slide text. Original slide design, images and animations are not preserved.",
        "PDF → Word": "Extracts selectable text into an editable Word file. Scanned pages may need OCR.",
        "Word → PDF": "Creates a simple PDF from Word paragraphs and tables; advanced formatting may differ.",
        "Excel → CSV": "Each worksheet becomes a CSV. Multiple worksheets are delivered in a ZIP.",
        "Image → PDF": "Convert an image into a PDF page.",
        "PDF → Images": "Exports every PDF page as PNG or JPEG in a ZIP.",
        "CSV → PDF": "Creates a table PDF using up to 500 rows.",
        "Excel → PDF": "Creates one simple table PDF per worksheet and packages them in a ZIP.",
    }
    st.markdown(f'<div class="panel">{notes[tool]}</div>', unsafe_allow_html=True)
    files = st.file_uploader("📂 Upload file(s)", type=TOOLS[tool],
                             accept_multiple_files=(batch and tool != "Excel → CSV"),
                             help="Maximum 200 MB per file. Conversion limits may also depend on hosting resources.")
    qlabel = st.select_slider("PDF → PowerPoint quality", ["Standard", "High", "Very high"], value="High")
    scale = {"Standard": 1.2, "High": 1.8, "Very high": 2.4}[qlabel]
    c1, c2 = st.columns(2)
    with c1:
        dpi = st.select_slider("PDF image DPI", [96, 150, 200, 300], value=150)
    with c2:
        img_fmt = st.selectbox("Image output format", ["PNG", "JPEG"])
    if files and show_details:
        st.markdown("#### Uploaded files")
        for f in files:
            st.write(f"**{f.name}** · {f.size / 1024 / 1024:.2f} MB")
    if files:
        total = sum(f.size for f in files)
        st.caption(f"{len(files)} file(s) selected · {total/1024/1024:.2f} MB total")
    do_convert = st.button("✨ Convert files", type="primary", use_container_width=True, disabled=not bool(files))
    b1, b2 = st.columns(2)
    with b1:
        if st.button("🧹 Clear results", use_container_width=True):
            st.session_state.results = []
            st.rerun()
    with b2:
        if st.button("Clear history", use_container_width=True):
            st.session_state.history = []
            st.rerun()

    if do_convert and files:
        outputs, errors = [], []
        bar = st.progress(0, text="Starting conversion...")
        for i, f in enumerate(files):
            try:
                if f.size == 0:
                    raise ValueError("This file is empty.")
                if f.size > 200 * 1024 * 1024:
                    raise ValueError("File exceeds the 200 MB limit.")
                payload = convert(tool, f, scale, dpi, img_fmt)
                if len(payload) == 3:
                    out_data, out_name, mime = payload
                else:
                    out_data, out_name, mime = payload
                outputs.append({"data": out_data, "name": out_name, "mime": mime, "source": f.name})
                st.session_state.history.append({"tool": tool, "source": f.name, "status": "Success"})
            except Exception as e:
                errors.append((f.name, str(e)))
                st.session_state.history.append({"tool": tool, "source": f.name, "status": "Failed"})
            bar.progress((i+1)/len(files), text=f"Processed {i+1} of {len(files)}")
        bar.empty()
        st.session_state.results = outputs
        if outputs:
            st.success(f"Converted {len(outputs)} file(s) successfully.")
        for name, error in errors:
            st.error(f"{name}: {error}")

    if st.session_state.results:
        st.markdown("#### ⬇️ Download results")
        for i, item in enumerate(st.session_state.results):
            ca, cb = st.columns([3, 1])
            with ca:
                st.write(f"**{item['name']}**")
                st.caption(f"From {item['source']} · {len(item['data'])/1024:.1f} KB")
            with cb:
                st.download_button("Download", item["data"], file_name=item["name"], mime=item["mime"],
                                   key=f"download_{i}_{item['name']}")
        if len(st.session_state.results) > 1:
            zbuf = io.BytesIO()
            with zipfile.ZipFile(zbuf, "w", zipfile.ZIP_DEFLATED) as z:
                for item in st.session_state.results:
                    z.writestr(item["name"], item["data"])
            st.download_button("📦 Download all as ZIP", zbuf.getvalue(), "filemorph_results.zip",
                               "application/zip", use_container_width=True)

with side:
    st.markdown('<div class="panel"><h3>📘 Quick guide</h3>', unsafe_allow_html=True)
    for n, title in enumerate(["Upload your files", "Choose quality/settings", "Convert and check errors", "Download results"], 1):
        st.markdown(f"**🔵 {n}. {title}**")
    st.markdown("</div>", unsafe_allow_html=True)
    st.markdown('<div class="panel"><h3>🛡️ Privacy</h3><p>Files are processed in the active app session. If hosted online, files pass through the hosting provider. Do not upload confidential documents to a public or untrusted app.</p></div>', unsafe_allow_html=True)
    st.markdown('<div class="panel"><h3>ℹ️ Notes</h3><p>Office-to-PDF conversions use simplified rendering to work without LibreOffice on Streamlit hosting. Original formatting, images, charts, animations and fonts may not be preserved.</p></div>', unsafe_allow_html=True)

with st.expander("🕘 Session conversion history"):
    if st.session_state.history:
        st.dataframe(pd.DataFrame(st.session_state.history), use_container_width=True, hide_index=True)
    else:
        st.caption("No conversions yet.")

st.caption("FileMorph AI · Streamlit file converter · Review converted files before use.")
